"""Fetcher de la OECD Affordable Housing Database (un fichero Excel por indicador).

Documentación: https://www.oecd.org/en/data/datasets/oecd-affordable-housing-database.html
Ficheros: https://webfs.oecd.org/els-com/Affordable_Housing_Database/ (listado del servidor,
verificado el 2026-10-10). Esta base no está en OECD Data Explorer: cada indicador (HC1.2, HM1.3,
PH4.2...) es un `.xlsx` con hojas de figuras y tablas anexas. Se leen dos formatos de hoja:

- **Tabla anual** (hojas de anexo, p. ej. `HM1.3.A1`): una fila de cabecera con los años y, por país,
  una o varias filas de categorías. El país va en la primera columna, en su propia fila o en la primera
  fila del bloque. `categoria` elige la fila del bloque; si es una lista, `valor` es la suma (vacío si
  falta alguna). Las categorías se reconocen porque se repiten en todos los países.
- **Pares valor/año** (p. ej. `Table PH4.2.A1`): una columna de valor (`columna`) seguida de su año
  («Year1», «Year2»...). Cada par da una observación por país.

Ejemplo real (hogares en vivienda en propiedad, con o sin hipoteca, % del total)::

    from lib import oecd_ahd

    df, fuentes = oecd_ahd.get(
        "HM1-3-Housing-tenures", "HM1.3.A1",
        categoria=["Own outright", "Owner with mortgage"],
        paises={"Spain": "espana", "United States": "eeuu", "Canada": "canada"},
    )

Ejemplo real (parque de vivienda social, % del total de viviendas, hacia 2010 y hacia 2022)::

    df, fuentes = oecd_ahd.get(
        "PH4-2-Social-rental-housing-stock", "Table PH4.2.A1",
        columna="% of total housing stock", paises={"Spain": "espana", "Canada": "canada"},
    )

Los nombres de país se comparan sin espacios sobrantes ni llamadas a nota («Spain (8)» → «Spain»).
«..», «-» y las celdas vacías son dato ausente. `escala` multiplica los valores (p. ej. 100 si la hoja
da proporciones en vez de porcentajes). Las medias de la OCDE solo aparecen en las hojas de figura y sin
año exacto, así que este fetcher no las lee. Los países sin datos no aparecen en el resultado ni en
`fuentes`.
"""

from __future__ import annotations

import io
import math
import re
from datetime import date
from typing import Mapping, Sequence

import openpyxl
import pandas as pd
import requests

from ._http import pedir
from .tipos import FuenteMeta, validar_df

BASE = "https://webfs.oecd.org/els-com/Affordable_Housing_Database"
ORGANISMO = "OECD"
AUSENTE = {"", "..", "...", "-", "–", "n.a.", "na", "x"}
_NOTA = re.compile(r"\s*\(\d+(?:\s*,\s*\d+)*\)\s*$")


class OecdAhdError(RuntimeError):
    """Error al descargar o leer un fichero de la OECD Affordable Housing Database."""


def _texto(v: object) -> str:
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


def _clave(v: object) -> str:
    """Normaliza un nombre de país o categoría para compararlo."""
    return _NOTA.sub("", _texto(v)).rstrip("*").strip().casefold()


def _num(v: object) -> float:
    if isinstance(v, bool):
        return math.nan
    if isinstance(v, (int, float)):
        return float(v)
    t = _texto(v)
    if t.lower() in AUSENTE:
        return math.nan
    try:
        return float(t)
    except ValueError:
        return math.nan


def _anio(v: object) -> int | None:
    t = _texto(v)
    if isinstance(v, float) and v.is_integer():
        t = str(int(v))
    return int(t) if re.fullmatch(r"(19|20)\d\d", t) else None


def leer_libro(contenido: bytes, fichero: str = "") -> openpyxl.Workbook:
    """Abre el `.xlsx` descargado; si no lo es (p. ej. una página HTML), lanza `OecdAhdError`."""
    try:
        return openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:  # zipfile.BadZipFile, KeyError, InvalidFileException...
        raise OecdAhdError(f"{fichero}: la respuesta no es un fichero Excel válido ({e})") from e


def _filas(libro: openpyxl.Workbook, hoja: str) -> list[tuple]:
    if hoja not in libro.sheetnames:
        raise ValueError(f"hoja {hoja!r} inexistente; hojas: {', '.join(libro.sheetnames)}")
    return list(libro[hoja].iter_rows(values_only=True))


def tabla_anual(filas: list[tuple], categoria: str | Sequence[str] | None) -> pd.DataFrame:
    """Lee una tabla de anexo. Devuelve `pais, periodo, valor` (`pais` tal como aparece, sin notas)."""
    cab = next((i for i, f in enumerate(filas) if sum(_anio(v) is not None for v in f) >= 3), None)
    if cab is None:
        raise ValueError("no hay una fila de cabecera con años (¿es una hoja de pares valor/año? usa `columna`)")
    anios = {j: a for j, v in enumerate(filas[cab]) if (a := _anio(v)) is not None}
    n_etiquetas = min(anios)
    cuerpo = filas[cab + 1 :]

    def etiquetas(f: tuple) -> list[str]:
        return [_texto(f[j]) if j < len(f) else "" for j in range(n_etiquetas)]

    # Con una sola columna de etiquetas, las categorías son las etiquetas que se repiten entre países.
    categorias: set[str] = set()
    if n_etiquetas == 1:
        cuenta: dict[str, int] = {}
        for f in cuerpo:
            if (e := _clave(etiquetas(f)[0])):
                cuenta[e] = cuenta.get(e, 0) + 1
        categorias = {e for e, n in cuenta.items() if n > 1}
    pais = ""
    obs: list[tuple[str, str, int, float]] = []
    for f in cuerpo:
        et = etiquetas(f)
        if n_etiquetas >= 2:
            pais = et[0] or pais
            cat = et[1]
        elif _clave(et[0]) in categorias:
            cat = et[0]
        else:
            pais, cat = et[0], ""
        for j, a in anios.items():
            v = _num(f[j]) if j < len(f) else math.nan
            if pais and not math.isnan(v):
                obs.append((pais, cat, a, v))
    tabla = pd.DataFrame(obs, columns=["pais", "categoria", "periodo", "valor"])

    disponibles = sorted({c for c in tabla["categoria"] if c})
    if categoria is None:
        if disponibles:
            raise ValueError(f"la hoja tiene categorías; elige una con `categoria`: {', '.join(disponibles)}")
        return tabla.drop(columns="categoria")
    pedidas = [categoria] if isinstance(categoria, str) else list(categoria)
    faltan = [c for c in pedidas if _clave(c) not in {_clave(d) for d in disponibles}]
    if faltan:
        raise ValueError(f"categorías inexistentes: {', '.join(faltan)}; disponibles: {', '.join(disponibles)}")
    tabla = tabla[tabla["categoria"].map(_clave).isin({_clave(c) for c in pedidas})]
    # Suma solo si están todas las categorías de ese país y año
    g = tabla.groupby(["pais", "periodo"])["valor"]
    suma = g.sum()[g.count() == len(pedidas)]
    return suma.reset_index()


def tabla_pares(filas: list[tuple], columna: str) -> pd.DataFrame:
    """Lee una tabla de pares valor/año. Devuelve `pais, periodo, valor`."""
    objetivo = _clave(columna)
    for i, f in enumerate(filas):
        cols = [j for j, v in enumerate(f) if _clave(v) == objetivo]
        if cols:
            break
    else:
        raise ValueError(f"columna {columna!r} no encontrada en la hoja")
    cab = filas[i]
    pares = []
    for j in cols:
        k = next((k for k in range(j + 1, len(cab)) if _clave(cab[k]).startswith("year")), None)
        if k is None:
            raise ValueError(f"la columna {columna!r} (posición {j + 1}) no tiene una columna «Year» a su derecha")
        pares.append((j, k))
    obs = []
    for f in filas[i + 1 :]:
        pais = _texto(f[0]) if f else ""
        for j, k in pares:
            v, a = (_num(f[j]), _anio(f[k])) if k < len(f) else (math.nan, None)
            if pais and a is not None and not math.isnan(v):
                obs.append((pais, a, v))
    tabla = pd.DataFrame(obs, columns=["pais", "periodo", "valor"])
    # Si dos pares caen en el mismo año, vale el de más a la derecha (el más reciente de la tabla).
    return tabla.drop_duplicates(["pais", "periodo"], keep="last")


def get(
    fichero: str,
    hoja: str,
    *,
    paises: Mapping[str, str],
    categoria: str | Sequence[str] | None = None,
    columna: str | None = None,
    escala: float = 1.0,
    desde: int = 1976,
    fecha_consulta: str | None = None,
    sesion: requests.Session | None = None,
    reintentos: int = 2,
) -> tuple[pd.DataFrame, list[FuenteMeta]]:
    """Descarga un indicador. Devuelve `(DataFrame serie/periodo/valor, fuentes)`.

    - `fichero`: nombre sin extensión, como aparece en el listado (p. ej. `HC1-2-Housing-costs-over-income`).
    - `paises`: nombre del país en el fichero → id de serie.
    - `categoria` (tabla anual) o `columna` (pares valor/año): uno de los dos según el formato de la hoja.
    """
    if categoria is not None and columna is not None:
        raise ValueError("usa `categoria` (tabla anual) o `columna` (pares valor/año), no ambas")
    url = f"{BASE}/{fichero}.xlsx"
    r = pedir(
        sesion or requests.Session(), url, None,
        error_cls=OecdAhdError, etiqueta=fichero, mensaje=lambda r: r.text[:200], reintentos=reintentos,
    )
    filas = _filas(leer_libro(r.content, fichero), hoja)
    tabla = tabla_pares(filas, columna) if columna is not None else tabla_anual(filas, categoria)

    serie_de = {_clave(k): v for k, v in paises.items()}
    tabla = tabla.assign(serie=tabla["pais"].map(_clave).map(serie_de)).dropna(subset=["serie"])
    tabla = tabla[tabla["periodo"] >= desde]
    df = validar_df(
        tabla.assign(valor=tabla["valor"] * escala)[["serie", "periodo", "valor"]].reset_index(drop=True)
    )

    detalle = columna if columna is not None else " + ".join([categoria] if isinstance(categoria, str) else categoria or [])
    codigo = f"Affordable Housing Database, {fichero} ({hoja}{': ' + detalle if detalle else ''})"
    hoy = fecha_consulta or date.today().isoformat()
    fuentes = [FuenteMeta(s, ORGANISMO, codigo, url, hoy) for s in df["serie"].unique()]
    return df, fuentes
