"""Fetcher del Servicio Público de Empleo Estatal (SEPE): informe mensual de prestaciones por desempleo.

Página: https://www.sepe.es/HomeSepe/que-es-el-sepe/estadisticas/estadisticas-prestaciones/informe-prestaciones.html
No hay API: cada mes se publica un Excel («Informe-AAAAMM.xlsx», con variantes en el nombre) con las tablas de
beneficiarios (BP-*) y de tasa de cobertura (TC-*). El fetcher lee la página, busca el enlace del informe de
diciembre de cada año y extrae las medias anuales. Mecanismo verificado el 2026-10-16; la página enlaza los informes
desde 2017.

- `get_cobertura(anios)`: tasa de cobertura bruta del total de beneficiarios (hoja TC-1.1a, fila «TOTAL»), media
  anual (columna «Media anual» o, si falta, media de los doce meses). Serie `cobertura`.
- `get_beneficiarios(anios)`: beneficiarios de prestaciones por desempleo, total (`beneficiarios`) y extranjeros
  (`beneficiarios_extranjeros`), media anual (hojas BP-1.1 y BP-1.4, bloque «Media a: dic», fila «Total» de ambos sexos).

Ejemplo real::

    from lib import sepe

    df, fuentes = sepe.get_cobertura(range(2017, 2026))

Los informes pesan unos 12 MB; con `ESPANAENDATOS_CACHE_HTTP` (lib/_http.py) no se vuelven a descargar.
"""

from __future__ import annotations

import io
import re
from datetime import date
from typing import Iterable

import openpyxl
import pandas as pd
import requests

from ._http import pedir
from .tipos import FuenteMeta, validar_df

ORGANISMO = "SEPE"
BASE = "https://www.sepe.es"
PAGINA = f"{BASE}/HomeSepe/que-es-el-sepe/estadisticas/estadisticas-prestaciones/informe-prestaciones.html"
CABECERA = {"User-Agent": "Mozilla/5.0 (compatible; EspanaenDatos)"}


class SepeError(RuntimeError):
    """Error de la web del SEPE o informe con un formato no reconocido."""


def _get(sesion: requests.Session, url: str, etiqueta: str) -> requests.Response:
    return pedir(sesion, url, None, error_cls=SepeError, etiqueta=etiqueta, mensaje=lambda r: r.text[:200],
                 headers=CABECERA)


def informes(html: str) -> dict[tuple[int, int], str]:
    """{(año, mes): url} de los informes en Excel enlazados en la página."""
    out = {}
    for href in re.findall(r'href="([^"]+\.xlsx)"', html, re.I):
        m = re.search(r"/datos/(\d{4})/[^/]*?(\d{4})(\d{2})\.xlsx$", href)
        if not m or m.group(1) != m.group(2):
            continue
        out.setdefault((int(m.group(2)), int(m.group(3))), href if href.startswith("http") else BASE + href)
    # Algún mes se publica con un nombre descriptivo (p. ej. «informe-de-prestaciones-por-desempleo-octubre-2019.xlsx»)
    meses = "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split()
    for href in re.findall(r'href="([^"]+\.xlsx)"', html, re.I):
        m = re.search(r"/datos/(\d{4})/[^/]*?(" + "|".join(meses) + r")-(\d{4})\.xlsx$", href, re.I)
        if m:
            out.setdefault((int(m.group(3)), meses.index(m.group(2).lower()) + 1), href if href.startswith("http") else BASE + href)
    return out


def _libro(sesion: requests.Session, url: str) -> openpyxl.Workbook:
    r = _get(sesion, url, url.rsplit("/", 1)[-1])
    try:
        return openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise SepeError(f"{url}: no es un Excel válido") from e


def leer_cobertura(libro: openpyxl.Workbook) -> float:
    """Media anual de la tasa de cobertura del total de beneficiarios (hoja TC-1.1a, fila «TOTAL»)."""
    if "TC-1.1a" not in libro.sheetnames:
        raise SepeError("falta la hoja TC-1.1a")
    filas = list(libro["TC-1.1a"].iter_rows(values_only=True))
    cab = next((f for f in filas if f and "Enero" in [str(x or "").strip() for x in f]), None)
    if cab is None:
        raise SepeError("TC-1.1a: falta la cabecera de meses")
    nombres = [str(x or "").strip() for x in cab]
    i_ene = nombres.index("Enero")
    i_media = nombres.index("Media anual") if "Media anual" in nombres else None
    for f in filas:
        if any(str(x or "").strip().upper() == "TOTAL" for x in f[: i_ene]):
            if i_media is not None and isinstance(f[i_media], (int, float)):
                return float(f[i_media])
            meses = [x for x in f[i_ene: i_ene + 12] if isinstance(x, (int, float))]
            if len(meses) != 12:
                raise SepeError("TC-1.1a: el año no tiene los doce meses")
            return sum(meses) / 12
    raise SepeError("TC-1.1a: falta la fila TOTAL")


def leer_beneficiarios(libro: openpyxl.Workbook, hoja: str) -> float:
    """Media anual de beneficiarios (bloque derecho «Media a:», fila «Total» de ambos sexos, columna TOTAL)."""
    if hoja not in libro.sheetnames:
        raise SepeError(f"falta la hoja {hoja}")
    filas = list(libro[hoja].iter_rows(values_only=True))
    cab = next((f for f in filas if f and [str(x or "").strip() for x in f].count("Edad y sexo") == 2), None)
    if cab is None:
        raise SepeError(f"{hoja}: falta la cabecera de los dos bloques")
    nombres = [str(x or "").strip() for x in cab]
    i_edad = [i for i, x in enumerate(nombres) if x == "Edad y sexo"][1]  # bloque de la media anual
    i_total = i_edad + 1
    for f in filas:
        if str(f[i_edad] or "").strip() == "Total" and isinstance(f[i_total], (int, float)):
            return float(f[i_total])  # la primera fila «Total» es la de ambos sexos
    raise SepeError(f"{hoja}: falta la fila Total")


def _diciembres(anios: Iterable[int], sesion: requests.Session) -> dict[int, openpyxl.Workbook]:
    disponibles = informes(_get(sesion, PAGINA, "informe de prestaciones").text)
    return {a: _libro(sesion, disponibles[(a, 12)]) for a in sorted(anios) if (a, 12) in disponibles}


def get_cobertura(anios: Iterable[int], *, fecha_consulta: str | None = None,
                  sesion: requests.Session | None = None) -> tuple[pd.DataFrame, list[FuenteMeta]]:
    """Tasa de cobertura bruta del sistema de protección por desempleo, media anual (serie `cobertura`)."""
    libros = _diciembres(anios, sesion or requests.Session())
    if not libros:
        raise SepeError("ningún informe de diciembre en los años pedidos")
    df = validar_df(pd.DataFrame([("cobertura", a, leer_cobertura(l)) for a, l in libros.items()],
                                 columns=["serie", "periodo", "valor"]))
    hoy = fecha_consulta or date.today().isoformat()
    return df, [FuenteMeta("cobertura", ORGANISMO, "Informe de prestaciones por desempleo de diciembre, TC-1.1a (media anual)", PAGINA, hoy)]


def get_beneficiarios(anios: Iterable[int], *, fecha_consulta: str | None = None,
                      sesion: requests.Session | None = None) -> tuple[pd.DataFrame, list[FuenteMeta]]:
    """Beneficiarios de prestaciones por desempleo, total y extranjeros, media anual."""
    libros = _diciembres(anios, sesion or requests.Session())
    if not libros:
        raise SepeError("ningún informe de diciembre en los años pedidos")
    filas = []
    for a, l in libros.items():
        filas += [("beneficiarios", a, leer_beneficiarios(l, "BP-1.1")),
                  ("beneficiarios_extranjeros", a, leer_beneficiarios(l, "BP-1.4"))]
    df = validar_df(pd.DataFrame(filas, columns=["serie", "periodo", "valor"]))
    hoy = fecha_consulta or date.today().isoformat()
    fuentes = [FuenteMeta(sr, ORGANISMO, f"Informe de prestaciones por desempleo de diciembre, {h} (media anual)", PAGINA, hoy)
               for sr, h in (("beneficiarios", "BP-1.1"), ("beneficiarios_extranjeros", "BP-1.4"))]
    return df, fuentes
