"""Fetcher de la brecha del IVA de la Comisión Europea (DG TAXUD), estudio anual «VAT gap in Europe».

Página del estudio: https://taxation-customs.ec.europa.eu/taxation/vat/fight-against-vat-fraud/vat-gap_en
Datos: el «data hub» de DG TAXUD enlazado desde esa página («Access the full data hub»), que funciona
sobre la misma plataforma que la base de datos de Eurostat y tiene una API SDMX 2.1 (aquí, en
formato SDMX-CSV). Dataflow `VAP_COMPL_GAP` («VAT Compliance Gap – How much VAT goes uncollected»),
con la edición más reciente del estudio, que `edicion()` lee de los metadatos del dataflow.
Verificado el 2026-10-11: edición «VAT Gap in Europe – report 2025» (actualización del 2025-12-10),
años 2013-2024. De esta edición solo son las estimaciones completas (2019-2023) y la rápida (2024,
flag `ep`): las cifras de 2013-2018 son las de ediciones anteriores sin ajustar (en 2016-2018
coinciden, para todos los países, con las del anexo de la edición de 2024). La serie de la edición para
los años anteriores a las estimaciones completas es la retropolada del anexo estadístico del informe
(tabla 205, desde 2000 y solo en %), que lee `get_todas()` (ver abajo).

Clave `A.{indicador}.{unidad}.{GEO}`:

- `INDIC_VAT_COMPL_GAP` (`indicador`): `VAT_COMPL_GAP` (brecha de cumplimiento), `VTTL` (obligación
  tributaria teórica total de IVA, *VAT total tax liability*), `VAT_REVENUE` (recaudación) y los
  componentes de la VTTL `LIAB_HH`, `LIAB_NPISH_GOV`, `LIAB_IC`, `LIAB_GFCF` y `LIAB_NET_ADJ`.
- `UNIT` (`unidad`): `PC_VTTL` (% de la VTTL), `MIO_EUR` (millones de euros) y `MIO_NAC` (millones
  de moneda nacional).
- `GEO`: Estados miembros con código Eurostat (`ES`, `DE`...; Grecia es `EL`), `EU26` (UE sin
  Luxemburgo), `UK` (hasta 2019) y países candidatos.

`OBS_FLAG`: `ep` («estimated, provisional»: las estimaciones rápidas del último año, con un método
simplificado) y `u` («low reliability»: Luxemburgo y tres países candidatos). Las observaciones con
un flag de `excluir_flags` se descartan (por defecto, `ep`).

Ejemplo real (brecha de cumplimiento del IVA de España, % de la VTTL)::

    from lib import brecha_iva

    df, fuentes = brecha_iva.get("VAT_COMPL_GAP", "PC_VTTL", geos={"ES": "espana"})
    brecha_iva.edicion()  # 'VAT Gap in Europe – report 2025'

Ejemplo real (brecha y VTTL en millones de euros, para agregar países)::

    geos = {"ES": "ES", "FR": "FR", "IT": "IT"}
    brecha, _ = brecha_iva.get("VAT_COMPL_GAP", "MIO_EUR", geos=geos)
    vttl, _ = brecha_iva.get("VTTL", "MIO_EUR", geos=geos)

Ojo con `EU26` en la edición 2025: su `PC_VTTL` es la media «EU (average)» de la tabla 205 del
informe: retropolada en 2013-2018 y, en 2019, 10,6, distinta del cociente de sus propias cifras en
euros (11,15) y de la cifra del texto del informe (11,1). En `MIO_EUR` falta la brecha de 2023. Para
la media de la UE en los años de las estimaciones completas, suma `VAT_COMPL_GAP` y `VTTL` en
`MIO_EUR` de los Estados miembros sin Luxemburgo (así lo define el informe, recuadro 2; en 2020-2023
coincide con `EU26` en `PC_VTTL`).

Anexo estadístico (`get_todas`, `leer_todas`): el `.xlsx` enlazado desde la página del estudio
(«Download underlying data»), `URL_ANEXO`, de la edición `EDICION_ANEXO`; la URL cambia con cada
edición, así que el notebook comprueba que coincida con `edicion()`. La hoja `Table 205` («VAT
compliance gap – all available estimates (%)») da la brecha de cada país y las medias «EU (average)»
(ponderada por la VTTL, sin Luxemburgo) y «EU (median)», en proporción (0,064 = 6,4 %), desde 2000. La
fila bajo los años marca el tipo de estimación de cada año: `Backcasted series` (`retropolada`: las de
ediciones anteriores ajustadas al nivel de las completas, sin recalcular el modelo), `Full estimates`
(`completa`) y `Forecast` (`rapida`). «.» es dato ausente. Las filas se nombran por su primera celda
sin asteriscos de nota («EU (average)*» → «EU (average)», «LU**» → «LU»)::

    df, fuentes, tipos = brecha_iva.get_todas({"ES": "espana", "EU (average)": "ue27"})
    tipos[2018], tipos[2019]  # ('retropolada', 'completa')

Los códigos inexistentes y los periodos sin datos devuelven HTTP 400/404 con un mensaje SOAP, que se
propaga como `BrechaIvaError`. Los geos válidos sin datos (p. ej. `HR` en 2013) no aparecen en el
DataFrame ni en `fuentes`: el notebook los completa con `null` (METODOLOGIA §2.4).
"""

from __future__ import annotations

import html
import io
import math
import re
from datetime import date, datetime
from typing import Mapping, Sequence

import xml.etree.ElementTree as ET

import openpyxl
import pandas as pd
import requests

from ._http import pedir
from .tipos import COLUMNAS, FuenteMeta, validar_df

BASE = "https://webgate.ec.europa.eu/taxation_customs/redisstat/api/dissemination/sdmx/2.1/data"
DATAFLOW = "VAP_COMPL_GAP"
URL_METADATOS = BASE.removesuffix("/data") + f"/dataflow/TAXUD/{DATAFLOW}/latest"
URL_DATABROWSER = (
    "https://webgate.ec.europa.eu/taxation_customs/redisstat/databrowser/view/VAP_COMPL_GAP/default/table?lang=en"
)
URL_ESTUDIO = "https://taxation-customs.ec.europa.eu/taxation/vat/fight-against-vat-fraud/vat-gap_en"
ORGANISMO = "Comisión Europea (DG TAXUD)"
# Anexo estadístico (xlsx) de la edición indicada: enlace «Download underlying data» de la página del estudio
URL_ANEXO = (
    "https://taxation-customs.ec.europa.eu/document/download/b3045343-a25a-4fcf-97c8-b27c413b95ad_en"
    "?filename=VAT-GAP-2025-tables.xlsx"
)
EDICION_ANEXO = "VAT Gap in Europe – report 2025"
HOJA_TODAS = "Table 205"
TITULO_TODAS = "all available estimates"
TIPOS_ESTIMACION = {"Backcasted series": "retropolada", "Full estimates": "completa", "Forecast": "rapida"}
FLAGS_EXCLUIDOS = ("ep",)
COLUMNAS_CSV = ("GEO", "TIME_PERIOD", "OBS_VALUE", "OBS_FLAG")

_CODIGO = re.compile(r"^[A-Z0-9_]+$")
_FAULT = re.compile(r"<faultstring>(.*?)</faultstring>", re.S)
_COMILLAS = re.compile(r"[“\"](.+?)[”\"]")
_NS = {"c": "http://www.sdmx.org/resources/sdmxml/schemas/v2_1/common"}


class BrechaIvaError(RuntimeError):
    """Error devuelto por la API del data hub de DG TAXUD o fallo de red persistente."""


def _mensaje(texto: str) -> str:
    """Mensaje legible de una respuesta de error (SOAP `faultstring` o el principio del cuerpo)."""
    m = _FAULT.search(texto)
    return html.unescape(m.group(1)).strip() if m else texto[:200]


def leer_csv(texto: str, etiqueta: str = DATAFLOW) -> pd.DataFrame:
    """Lee la respuesta SDMX-CSV tal cual (todas las columnas como texto, celdas vacías = "")."""
    if not texto.strip():
        raise BrechaIvaError(f"{etiqueta}: respuesta vacía")
    if texto.lstrip().startswith("<"):
        raise BrechaIvaError(f"{etiqueta}: la respuesta no es SDMX-CSV: {_mensaje(texto)}")
    raw = pd.read_csv(io.StringIO(texto), dtype=str, keep_default_na=False)
    faltan = [c for c in COLUMNAS_CSV if c not in raw.columns]
    if faltan:
        raise BrechaIvaError(f"{etiqueta}: la respuesta no es SDMX-CSV (faltan {', '.join(faltan)})")
    return raw


def actualizacion(texto: str) -> str | None:
    """Fecha ISO de la columna `LAST UPDATE` (p. ej. «10/12/25 19:41:26» → «2025-12-10»), o None."""
    try:
        raw = leer_csv(texto)
        valor = raw["LAST UPDATE"].iloc[0]
        return datetime.strptime(valor.strip(), "%d/%m/%y %H:%M:%S").date().isoformat()
    except (BrechaIvaError, KeyError, IndexError, ValueError):
        return None


def csv_a_df(
    texto: str,
    geos: Mapping[str, str],
    *,
    excluir_flags: Sequence[str] = FLAGS_EXCLUIDOS,
    etiqueta: str = DATAFLOW,
) -> pd.DataFrame:
    """Convierte SDMX-CSV al DataFrame estándar `serie, periodo, valor` (solo los geos de `geos`)."""
    raw = leer_csv(texto, etiqueta)
    raw = raw[raw["GEO"].isin(geos)]
    if not raw["TIME_PERIOD"].str.fullmatch(r"\d{4}").all():
        raise ValueError(f"{etiqueta}: periodo no anual; solo se admiten años")
    dup = raw.duplicated(["GEO", "TIME_PERIOD"], keep=False)
    if dup.any():
        sobran = [c for c in raw.columns if c not in COLUMNAS_CSV and raw.loc[dup, c].nunique() > 1]
        raise ValueError(f"{etiqueta}: varias observaciones por (geo, año); fija las dimensiones: {', '.join(sobran)}")
    raw = raw[~raw["OBS_FLAG"].str.strip().isin(set(excluir_flags))]
    valores = raw["OBS_VALUE"].str.strip().replace("", math.nan)
    df = pd.DataFrame(
        {
            "serie": raw["GEO"].map(geos),
            "periodo": raw["TIME_PERIOD"].astype(int),
            "valor": pd.to_numeric(valores, errors="raise"),
        }
    )[COLUMNAS].dropna(subset=["valor"])
    return validar_df(df)


def get(
    indicador: str,
    unidad: str,
    *,
    geos: Mapping[str, str],
    desde: int = 1976,
    excluir_flags: Sequence[str] = FLAGS_EXCLUIDOS,
    fecha_consulta: str | None = None,
    sesion: requests.Session | None = None,
    reintentos: int = 2,
) -> tuple[pd.DataFrame, list[FuenteMeta]]:
    """Descarga un indicador anual del dataflow `VAP_COMPL_GAP`. Devuelve `(DataFrame serie/periodo/valor, fuentes)`.

    - `indicador`: código de `INDIC_VAT_COMPL_GAP` (p. ej. `VAT_COMPL_GAP`, `VTTL`).
    - `unidad`: código de `UNIT` (`PC_VTTL`, `MIO_EUR` o `MIO_NAC`).
    - `geos`: código GEO → id de serie (p. ej. `{"ES": "espana"}`).
    - `desde`: primer año (`startPeriod`); `excluir_flags`: flags de observación que se descartan.
    `dataset_codigo` de cada fuente incluye la fecha de actualización del dataflow (la edición del estudio).
    """
    geos = dict(geos)
    if not geos:
        raise ValueError("`geos` no puede estar vacío")
    for codigo in (indicador, unidad, *geos):
        if not isinstance(codigo, str) or not _CODIGO.match(codigo):
            raise ValueError(f"código no válido: {codigo!r} (solo mayúsculas, cifras y _)")
    url = f"{BASE}/{DATAFLOW}/A.{indicador}.{unidad}.{'+'.join(geos)}"
    r = pedir(
        sesion or requests.Session(),
        url,
        {"format": "SDMX-CSV", "startPeriod": str(desde)},
        error_cls=BrechaIvaError,
        etiqueta=DATAFLOW,
        mensaje=lambda r: _mensaje(r.text),
        reintentos=reintentos,
    )
    df = csv_a_df(r.text, geos, excluir_flags=excluir_flags)
    df = df[df["periodo"] >= desde].reset_index(drop=True)
    act = actualizacion(r.text)
    codigo = f"{DATAFLOW} ({indicador}, {unidad}{f'; actualización del {act}' if act else ''})"
    hoy = fecha_consulta or date.today().isoformat()
    fuentes = [FuenteMeta(s, ORGANISMO, codigo, URL_DATABROWSER, hoy) for s in df["serie"].unique()]
    return df, fuentes


def edicion(*, sesion: requests.Session | None = None, reintentos: int = 2) -> str:
    """Edición del estudio que publica el dataflow (p. ej. «VAT Gap in Europe – report 2025»).

    Sale de la anotación `SUBTITLE` de los metadatos del dataflow («Data from “VAT Gap in Europe –
    report 2025”»): sirve para nombrar la edición en las fuentes y las notas, y para detectar una nueva.
    """
    r = pedir(
        sesion or requests.Session(),
        URL_METADATOS,
        None,
        error_cls=BrechaIvaError,
        etiqueta=f"{DATAFLOW} (metadatos)",
        mensaje=lambda r: _mensaje(r.text),
        reintentos=reintentos,
    )
    try:
        raiz = ET.fromstring(r.content)  # bytes: el XML declara su codificación
    except ET.ParseError as e:
        raise BrechaIvaError(f"{DATAFLOW} (metadatos): la respuesta no es XML: {_mensaje(r.text)}") from e
    for anotacion in raiz.iter(f"{{{_NS['c']}}}Annotation"):
        if anotacion.findtext("c:AnnotationType", namespaces=_NS) == "SUBTITLE":
            texto = (anotacion.findtext("c:AnnotationText", namespaces=_NS) or "").strip()
            m = _COMILLAS.search(texto)
            if texto:
                return m.group(1).strip() if m else texto
    raise BrechaIvaError(f"{DATAFLOW} (metadatos): sin la anotación SUBTITLE con la edición del estudio")


def _etiqueta(v: object) -> str:
    """Texto de una celda sin espacios duros ni asteriscos de nota («EU (average)*» → «EU (average)»)."""
    return "" if v is None else str(v).replace("\xa0", " ").strip().rstrip("*").strip()


def leer_todas(
    contenido: bytes, filas: Mapping[str, str], *, hoja: str = HOJA_TODAS
) -> tuple[pd.DataFrame, dict[int, str]]:
    """Lee la tabla «all available estimates» del anexo (`.xlsx` en bytes).

    `filas`: etiqueta de la primera celda → id de serie (p. ej. `{"ES": "espana", "EU (average)": "ue27"}`).
    Devuelve `(DataFrame serie/periodo/valor en %, {año: "retropolada" | "completa" | "rapida"})`.
    """
    etiqueta = f"anexo ({hoja})"
    filas = dict(filas)
    if not filas:
        raise ValueError("`filas` no puede estar vacío")
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as e:  # zipfile.BadZipFile, KeyError, InvalidFileException...
        raise BrechaIvaError(f"{etiqueta}: la respuesta no es un fichero Excel válido ({e})") from e
    if hoja not in libro.sheetnames:
        raise BrechaIvaError(f"{etiqueta}: hoja inexistente; hojas: {', '.join(libro.sheetnames)}")
    rows = [tuple(r) for r in libro[hoja].iter_rows(values_only=True)]
    titulo = next((_etiqueta(c) for r in rows[:3] for c in r if isinstance(c, str) and c.strip()), "")
    if TITULO_TODAS not in titulo.lower():
        raise BrechaIvaError(f"{etiqueta}: no es la tabla «{TITULO_TODAS}» (título: {titulo!r})")

    def es_anio(x: object) -> bool:
        return isinstance(x, int) and not isinstance(x, bool) and 1990 <= x <= 2100

    i_cab = next((i for i, r in enumerate(rows) if sum(map(es_anio, r)) >= 5), None)
    if i_cab is None or i_cab + 1 >= len(rows):
        raise BrechaIvaError(f"{etiqueta}: sin fila de años")
    anios = {j: x for j, x in enumerate(rows[i_cab]) if es_anio(x)}
    # Tipo de estimación: fila bajo los años, en celdas combinadas (cada etiqueta vale hasta la siguiente)
    marcas, tipos, actual = rows[i_cab + 1], {}, None
    for j in sorted(anios):
        texto = _etiqueta(marcas[j]) if j < len(marcas) else ""
        if texto:
            if texto not in TIPOS_ESTIMACION:
                raise BrechaIvaError(f"{etiqueta}: tipo de estimación desconocido {texto!r}")
            actual = TIPOS_ESTIMACION[texto]
        if actual is None:
            raise BrechaIvaError(f"{etiqueta}: el año {anios[j]} no tiene tipo de estimación")
        tipos[anios[j]] = actual

    encontradas: dict[str, tuple] = {}
    for r in rows[i_cab + 2 :]:
        clave = _etiqueta(r[0]) if r else ""
        if clave in filas:
            if clave in encontradas:
                raise BrechaIvaError(f"{etiqueta}: fila repetida {clave!r}")
            encontradas[clave] = r
    faltan = [c for c in filas if c not in encontradas]
    if faltan:
        raise BrechaIvaError(f"{etiqueta}: filas inexistentes: {', '.join(faltan)}")
    registros = []
    for clave, r in encontradas.items():
        for j, anio in anios.items():
            v = r[j] if j < len(r) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                if abs(v) > 1:
                    raise BrechaIvaError(f"{etiqueta}: {clave} {anio}: {v} no es una proporción")
                registros.append((filas[clave], anio, round(v * 100, 6)))
            elif _etiqueta(v) not in {"", "."}:
                raise BrechaIvaError(f"{etiqueta}: {clave} {anio}: valor no numérico {v!r}")
    df = pd.DataFrame(registros, columns=COLUMNAS)
    return validar_df(df), tipos


def get_todas(
    filas: Mapping[str, str],
    *,
    desde: int = 1976,
    tipos: Sequence[str] = ("retropolada", "completa"),
    url: str = URL_ANEXO,
    fecha_consulta: str | None = None,
    sesion: requests.Session | None = None,
    reintentos: int = 2,
) -> tuple[pd.DataFrame, list[FuenteMeta], dict[int, str]]:
    """Descarga el anexo estadístico y lee la tabla 205 («all available estimates», %).

    Devuelve `(DataFrame serie/periodo/valor, fuentes, {año: tipo de estimación})`. Solo se conservan
    los años con un tipo de `tipos` (por defecto, sin las estimaciones rápidas) y desde `desde`.
    """
    desconocidos = set(tipos) - set(TIPOS_ESTIMACION.values())
    if desconocidos:
        raise ValueError(f"tipos desconocidos: {', '.join(sorted(desconocidos))}")
    r = pedir(
        sesion or requests.Session(),
        url,
        None,
        error_cls=BrechaIvaError,
        etiqueta=f"anexo ({HOJA_TODAS})",
        mensaje=lambda r: r.text[:200],
        reintentos=reintentos,
    )
    df, tipos_anio = leer_todas(r.content, filas)
    df = df[df["periodo"].map(tipos_anio).isin(set(tipos)) & (df["periodo"] >= desde)].reset_index(drop=True)
    codigo = f"anexo estadístico de «{EDICION_ANEXO}», {HOJA_TODAS} (VAT compliance gap – all available estimates, %)"
    hoy = fecha_consulta or date.today().isoformat()
    fuentes = [FuenteMeta(s, ORGANISMO, codigo, url, hoy) for s in df["serie"].unique()]
    return df, fuentes, tipos_anio
